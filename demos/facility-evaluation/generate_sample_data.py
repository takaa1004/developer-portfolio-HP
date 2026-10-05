"""デモ用の架空データとルール表を生成する。

実在の事業所・人物とは関係ありません。乱数は固定しているので、何度実行しても同じデータになります。

    python generate_sample_data.py
"""

import random
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data" / "sample"
RULES_PATH = BASE_DIR / "rules" / "評価ルール.xlsx"

MONTHS = pd.period_range("2026-04", "2026-09", freq="M")

NAMES = [
    "さくら", "ひまわり", "あおば", "みどり", "つばさ", "ひかり", "こもれび", "やまびこ", "せせらぎ", "なごみ",
    "ほほえみ", "かがやき", "いずみ", "わかば", "はるかぜ", "そよかぜ", "あさひ", "ゆうひ", "すずらん", "たんぽぽ",
]
PREFIX = {"通所介護": "デイサービス", "訪問介護": "ヘルパーステーション", "居宅介護支援": "ケアプランセンター"}
TYPES = ["通所介護"] * 9 + ["訪問介護"] * 7 + ["居宅介護支援"] * 4
AREAS = ["北エリア", "中央エリア", "南エリア"]

# 事業所ごとの「実力」。1.0で基準値どおり。デモ用のシナリオを仕込んでいる。
PERFORMANCE = {
    "F004": 0.83,  # やや低迷 → 業務改善報告書の対象
    "F012": 0.68,  # 大きく低迷 → 業務改善報告書・再教育の対象
    "F019": 0.80,  # 開設1年未満。既存の基準では低いが、新規の基準では平均以上
}
OPENED = {
    "F007": "2026-06",  # 開設4ヶ月 → 参考評価
    "F019": "2025-11",  # 開設1年未満 → 新規の基準値を適用
}


def build_facilities(rng: random.Random) -> pd.DataFrame:
    rows = []
    shuffled = TYPES[:]
    rng.shuffle(shuffled)
    for i, (name, kind) in enumerate(zip(NAMES, shuffled), start=1):
        fid = f"F{i:03d}"
        # シナリオ用に種別を固定する
        kind = {"F004": "通所介護", "F012": "通所介護", "F018": "通所介護", "F007": "訪問介護"}.get(fid, kind)
        capacity = rng.choice([18, 20, 25, 30, 35]) if kind == "通所介護" else None
        opened = OPENED.get(fid, f"{rng.randint(2012, 2023)}-{rng.randint(1, 12):02d}")
        rows.append({
            "事業所ID": fid,
            "事業所名": f"{PREFIX[kind]}{name}",
            "事業種別": kind,
            "エリア": AREAS[(i - 1) % 3],
            "定員": capacity,
            "開設年月": opened,
        })
    return pd.DataFrame(rows)


def build_results(facilities: pd.DataFrame, rng: random.Random) -> pd.DataFrame:
    rows = []
    for f in facilities.itertuples():
        base = PERFORMANCE.get(f.事業所ID, rng.uniform(0.88, 1.15))
        opened = pd.Period(f.開設年月, freq="M")
        for m in MONTHS:
            if m < opened:
                continue
            if f.事業所ID == "F015" and str(m) == "2026-07":
                continue  # 1ヶ月分の実績が未提出 → 要確認
            factor = base * rng.uniform(0.95, 1.05)
            if (m - opened).n < 3:
                factor *= 0.7  # 開設直後は立ち上がり期間
            days = 26 if m.month in (5, 7, 8) else 25
            if f.事業種別 == "通所介護":
                staff = round(f.定員 * 0.3 + rng.uniform(-0.5, 0.5), 1)
                visits = round(f.定員 * days * min(0.8 * factor, 0.97))
                units = round(visits * rng.uniform(830, 870))
            elif f.事業種別 == "訪問介護":
                staff = round(rng.uniform(9, 11), 1)
                visits = round(550 * factor)
                units = round(visits * rng.uniform(410, 430))
            else:
                staff = float(rng.choice([3, 4, 5]))
                visits = round(staff * 35 * factor)
                units = round(visits * rng.uniform(1380, 1420))
            if f.事業所ID == "F018" and str(m) == "2026-08":
                visits *= 10  # 入力ミス(桁違い) → 稼働率100%超で要確認
            rows.append({
                "事業所ID": f.事業所ID,
                "年月": str(m),
                "延べ利用者数": visits,
                "算定単位数": units,
                "営業日数": days,
                "スタッフ数": staff,
            })
    return pd.DataFrame(rows)


def build_assignments(facilities: pd.DataFrame, rng: random.Random) -> pd.DataFrame:
    family = ["佐藤", "鈴木", "高橋", "田中", "伊藤", "渡辺", "山本", "中村", "小林", "加藤", "吉田", "山田",
              "佐々木", "山口", "松本", "井上", "木村", "林", "清水", "山崎", "森", "池田", "橋本", "阿部"]
    given = ["一郎", "花子", "健太", "美咲", "大輔", "陽子", "翔", "由美", "誠", "直子", "拓也", "恵",
             "隆", "真理", "浩二", "愛", "修", "綾", "剛", "智子", "亮", "舞", "博", "優子"]
    rows = []
    for i, f in enumerate(facilities.itertuples(), start=1):
        sid = f"S{i:03d}"
        assigned = max(pd.Period(f.開設年月, freq="M"), pd.Period("2024-04", freq="M") - rng.randint(0, 24))
        left = "2026-07" if f.事業所ID == "F003" else ""  # F003は期中に責任者が交代
        rows.append({"社員ID": sid, "氏名": f"{family[i - 1]} {given[i - 1]}", "役職": "事業所責任者",
                     "事業所ID": f.事業所ID, "エリア": f.エリア, "着任年月": str(assigned), "離任年月": left})
    rows.append({"社員ID": "S021", "氏名": "森 亮", "役職": "事業所責任者", "事業所ID": "F003",
                 "エリア": facilities.loc[2, "エリア"], "着任年月": "2026-08", "離任年月": ""})
    managers = [("A001", "池田 舞", "北エリア", "2022-04"), ("A002", "橋本 博", "中央エリア", "2023-10"),
                ("A003", "阿部 優子", "南エリア", "2026-06")]  # A003は期中着任
    for aid, name, area, assigned in managers:
        rows.append({"社員ID": aid, "氏名": name, "役職": "エリアマネージャー", "事業所ID": "",
                     "エリア": area, "着任年月": assigned, "離任年月": ""})
    return pd.DataFrame(rows)


def build_leaves() -> pd.DataFrame:
    return pd.DataFrame([
        {"社員ID": "S005", "休職開始年月": "2026-05", "休職終了年月": "2026-08"},  # 4ヶ月休職 → 対象外
        {"社員ID": "S009", "休職開始年月": "2026-09", "休職終了年月": ""},  # 1ヶ月のみ → 残り5ヶ月で評価
    ])


def build_rules() -> dict[str, pd.DataFrame]:
    settings = pd.DataFrame([
        {"項目": "評価期間(月数)", "値": 6, "説明": "評価対象月を含めて遡る月数"},
        {"項目": "新規判定(開設後の月数)", "値": 12, "説明": "開設からこの月数未満は「新規」の基準値を適用"},
        {"項目": "達成率の上限(%)", "値": 150, "説明": "1指標の達成率がこれを超えたら上限で打ち切る"},
    ])
    baselines = pd.DataFrame([
        ("通所介護", "既存", "稼働率", 80, 0.6),
        ("通所介護", "既存", "1人当たり単位数", 58000, 0.4),
        ("通所介護", "新規", "稼働率", 60, 0.6),
        ("通所介護", "新規", "1人当たり単位数", 45000, 0.4),
        ("訪問介護", "既存", "延べ利用者数", 550, 0.4),
        ("訪問介護", "既存", "1人当たり単位数", 23000, 0.6),
        ("訪問介護", "新規", "延べ利用者数", 400, 0.4),
        ("訪問介護", "新規", "1人当たり単位数", 17000, 0.6),
        ("居宅介護支援", "既存", "1人当たり利用者数", 35, 0.6),
        ("居宅介護支援", "既存", "1人当たり単位数", 49000, 0.4),
        ("居宅介護支援", "新規", "1人当たり利用者数", 25, 0.6),
        ("居宅介護支援", "新規", "1人当たり単位数", 35000, 0.4),
    ], columns=["事業種別", "開設区分", "指標", "基準値", "重み"])
    ranks = pd.DataFrame([
        {"ランク": "S", "スコア下限": 110}, {"ランク": "A", "スコア下限": 100}, {"ランク": "B", "スコア下限": 90},
        {"ランク": "C", "スコア下限": 80}, {"ランク": "D", "スコア下限": 0},
    ])
    exceptions = pd.DataFrame([
        ("EX01", "事業所", "開設経過月数_未満", 6, "参考評価", "開設6ヶ月未満の事業所は参考評価(順位・抽出の対象外)", "○"),
        ("EX02", "事業所", "欠損月数_以上", 1, "要確認", "評価期間内に実績が未提出の月がある", "○"),
        ("EX03", "事業所", "稼働率_超過", 100, "要確認", "稼働率が100%を超える月がある(入力ミスの可能性)", "○"),
        ("EX11", "対象者", "着任経過月数_未満", 3, "対象外", "着任3ヶ月未満は評価対象外", "○"),
        ("EX12", "対象者", "休職月数_以上", 3, "対象外", "評価期間内に3ヶ月以上休職している", "○"),
        ("EX13", "対象者", "評価月数_未満", 3, "要確認", "評価に使える月が3ヶ月未満", "○"),
    ], columns=["ルールID", "対象", "条件", "閾値", "処理", "説明", "有効"])
    extractions = pd.DataFrame([
        ("業務改善報告書対象", "平均スコア_未満", "85", "", "6ヶ月平均スコアが85未満", "○"),
        ("業務改善報告書対象", "連続ランク_以下", "C", "3", "直近3ヶ月連続でC以下", "○"),
        ("再教育対象", "平均スコア_未満", "75", "", "6ヶ月平均スコアが75未満", "○"),
        ("再教育対象", "当月ランク_以下", "D", "", "当月ランクがD", "○"),
    ], columns=["抽出区分", "条件", "閾値", "閾値2", "説明", "有効"])
    return {"設定": settings, "基準値": baselines, "評価ランク": ranks, "例外条件": exceptions, "抽出条件": extractions}


def write_rules(sheets: dict[str, pd.DataFrame]) -> None:
    RULES_PATH.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(RULES_PATH, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)

    # 担当者が編集しやすいように見た目と入力規則を整える
    wb = load_workbook(RULES_PATH)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) * 2 + 2
            ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 70)
        ws.freeze_panes = "A2"

    def add_list(ws, col: str, options: list[str]) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}200")
        ws.add_data_validation(dv)

    add_list(wb["基準値"], "B", ["既存", "新規"])
    add_list(wb["基準値"], "C", ["稼働率", "延べ利用者数", "算定単位数", "1人当たり単位数", "1人当たり利用者数"])
    add_list(wb["例外条件"], "B", ["事業所", "対象者"])
    add_list(wb["例外条件"], "C", ["開設経過月数_未満", "欠損月数_以上", "稼働率_超過",
                                  "着任経過月数_未満", "休職月数_以上", "評価月数_未満"])
    add_list(wb["例外条件"], "E", ["対象外", "要確認", "参考評価"])
    add_list(wb["例外条件"], "G", ["○", "×"])
    add_list(wb["抽出条件"], "B", ["平均スコア_未満", "当月ランク_以下", "連続ランク_以下"])
    add_list(wb["抽出条件"], "F", ["○", "×"])
    wb.save(RULES_PATH)


def main() -> None:
    rng = random.Random(20260930)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    facilities = build_facilities(rng)
    outputs = {
        "事業所マスタ.csv": facilities,
        "月次実績.csv": build_results(facilities, rng),
        "着任情報.csv": build_assignments(facilities, rng),
        "休職情報.csv": build_leaves(),
    }
    for name, df in outputs.items():
        # 社内システムからの出力を想定して Shift-JIS(cp932) で保存する
        df.to_csv(DATA_DIR / name, index=False, encoding="cp932")
        print(f"作成: {DATA_DIR / name} ({len(df)}行)")
    write_rules(build_rules())
    print(f"作成: {RULES_PATH}")


if __name__ == "__main__":
    main()
