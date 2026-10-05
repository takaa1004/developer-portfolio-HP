"""評価結果を Excel に出力する。"""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import EvaluationResult
from .loader import InputData
from .rules import Rules

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FONT = Font(bold=True, size=12, color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
RANK_FILL = {
    "S": PatternFill("solid", fgColor="C6EFCE"), "A": PatternFill("solid", fgColor="E2F0D9"),
    "C": PatternFill("solid", fgColor="FFF2CC"), "D": PatternFill("solid", fgColor="F8CBAD"),
}
JUDGMENT_FILL = {
    "参考評価": PatternFill("solid", fgColor="DDEBF7"), "要確認": PatternFill("solid", fgColor="FFE699"),
    "対象外": PatternFill("solid", fgColor="D9D9D9"),
}


def _top_bottom(df: pd.DataFrame, score_col: str, rank_col: str, n: int = 3):
    cols = ["事業所ID", "事業所名", "事業種別", "エリア", score_col, rank_col]
    target = df[df["判定"] == "通常"].dropna(subset=[score_col])
    return target.nlargest(n, score_col)[cols], target.nsmallest(n, score_col)[cols]


def write_report(result: EvaluationResult, data: InputData, rules: Rules, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    n = len(result.period)
    avg_col, avg_rank = f"{n}ヶ月平均スコア", f"{n}ヶ月平均ランク"
    fac, persons = result.facilities, result.persons
    period_label = f"{result.period[0]} 〜 {result.period[-1]}({n}ヶ月)"

    # 月次レポート用: 事業所×月のスコア推移
    trend = result.facility_monthly.pivot(index="事業所ID", columns="年月", values="スコア")
    trend.columns = [str(c) for c in trend.columns]
    trend = fac[["事業所ID", "事業所名", "判定"]].merge(trend.reset_index(), on="事業所ID", how="left")

    summary = [
        ("評価対象月", str(result.target_month)),
        ("評価期間", period_label),
        ("実行日時", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("事業所数", len(fac)),
        *[(f"　判定: {k}", int((fac["判定"] == k).sum())) for k in ["通常", "参考評価", "要確認"]],
        ("対象者数", len(persons)),
        *[(f"　判定: {k}", int((persons["判定"] == k).sum())) for k in ["通常", "参考評価", "要確認", "対象外"]],
        ("", ""),
        *[(f"{k}", f"{len(v)}名") for k, v in result.extracted.items()],
        ("要確認一覧", f"{len(result.review)}件"),
        ("", ""),
        ("次の作業", "「要確認一覧」の各行を確認し、確認者・確認結果・確認日を記入してください。"),
        ("", "すべての確認が済むまで、評価結果を公開しないでください。"),
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(summary, columns=["項目", "内容"]).to_excel(writer, sheet_name="サマリー", index=False)
        fac.to_excel(writer, sheet_name="事業所別評価", index=False)
        persons.to_excel(writer, sheet_name="対象者別評価", index=False)
        for name, df in result.extracted.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
        result.review.to_excel(writer, sheet_name="要確認一覧", index=False)

        sheet, row = "月次レポート", 0
        sections = [
            (f"当月({result.target_month})上位3事業所", _top_bottom(fac, "当月スコア", "当月ランク")[0]),
            (f"当月({result.target_month})下位3事業所", _top_bottom(fac, "当月スコア", "当月ランク")[1]),
            (f"{n}ヶ月平均 上位3事業所", _top_bottom(fac, avg_col, avg_rank)[0]),
            (f"{n}ヶ月平均 下位3事業所", _top_bottom(fac, avg_col, avg_rank)[1]),
            ("事業所別スコア推移", trend),
        ]
        section_rows = []
        for title, df in sections:
            section_rows.append((row + 1, title))
            df.to_excel(writer, sheet_name=sheet, index=False, startrow=row + 1)
            row += len(df) + 4

        result.percentiles.to_excel(writer, sheet_name="パーセンタイル", index=False)
        detail = result.facility_detail.copy()
        detail["年月"] = detail["年月"].astype(str)
        detail.to_excel(writer, sheet_name="算出根拠_事業所", index=False)
        pdetail = result.person_detail.copy()
        pdetail["年月"] = pdetail["年月"].astype(str)
        pdetail.to_excel(writer, sheet_name="算出根拠_対象者", index=False)

        log = pd.DataFrame(data.files + [rules.source])
        log.to_excel(writer, sheet_name="実行ログ", index=False, startrow=1)
        settings = pd.DataFrame([
            ("評価期間(月数)", rules.period_months), ("新規判定(開設後の月数)", rules.new_facility_months),
            ("達成率の上限(%)", rules.achievement_cap), ("有効な例外条件", len(rules.exceptions)),
            ("有効な抽出条件", len(rules.extractions)),
        ], columns=["設定", "値"])
        settings.to_excel(writer, sheet_name="実行ログ", index=False, startrow=len(log) + 5)

    _style(path, section_rows, len(data.files) + 1)
    return path


def _style(path: Path, section_rows: list[tuple[int, str]], log_rows: int) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        header_rows = {1}
        if ws.title == "月次レポート":
            header_rows = {r + 1 for r, _ in section_rows}
            for r, title in section_rows:
                ws.cell(row=r, column=1, value=title).font = SECTION_FONT
        elif ws.title == "実行ログ":
            header_rows = {2, log_rows + 6}
            ws.cell(row=1, column=1, value="入力ファイル(同じハッシュなら同じファイルで実行したことを確認できます)").font = SECTION_FONT
            ws.cell(row=log_rows + 5, column=1, value="適用した設定").font = SECTION_FONT

        headers = {}
        for r in header_rows:
            for cell in ws[r]:
                if cell.value is None:
                    continue
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                headers[(r, cell.column)] = str(cell.value)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.row in header_rows or cell.value is None:
                    continue
                cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
                header = _header_for(headers, cell)
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0" if abs(cell.value) >= 1000 else "0.0"
                if header and "ランク" in header and cell.value in RANK_FILL:
                    cell.fill = RANK_FILL[cell.value]
                if header == "判定" and cell.value in JUDGMENT_FILL:
                    cell.fill = JUDGMENT_FILL[cell.value]

        for col in range(1, ws.max_column + 1):
            values = [str(ws.cell(row=r, column=col).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
            width = max((sum(2 if ord(ch) > 255 else 1 for ch in v) for v in values), default=8) + 2
            ws.column_dimensions[get_column_letter(col)].width = min(max(width, 8), 60)
        if ws.title not in ("月次レポート", "実行ログ", "サマリー"):
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions

    summary = wb["サマリー"]
    summary.column_dimensions["B"].width = 70
    wb.save(path)


def _header_for(headers: dict, cell) -> str | None:
    """セルの上にある一番近い見出しを返す(1シートに表が複数ある場合に対応)。"""
    rows = [r for (r, c) in headers if c == cell.column and r < cell.row]
    return headers[(max(rows), cell.column)] if rows else None
