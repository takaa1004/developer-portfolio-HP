"""事業所評価の簡易デモ。

    python evaluate.py --month 2026-09

流れ: CSV読み込み → ルール表に従ってスコア算出 → 例外判定 → Excel出力
計算に AI や乱数は使わないので、同じ入力とルールからは必ず同じ結果になります。
"""

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
PRIORITY = {"通常": 0, "参考評価": 1, "要確認": 2}
RANK_FILL = {"S": "C6EFCE", "A": "E2F0D9", "C": "FFF2CC", "D": "F8CBAD"}
JUDGE_FILL = {"参考評価": "DDEBF7", "要確認": "FFE699"}


def read_csv(path: Path) -> pd.DataFrame:
    for enc in ("utf-8-sig", "cp932"):  # UTF-8 でも Shift-JIS でも読める
        try:
            return pd.read_csv(path, encoding=enc)
        except UnicodeDecodeError:
            pass
    sys.exit(f"[エラー] {path.name}: 文字コードを判別できません")


def read_rule_sheet(path: Path, name: str) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=name)
    df.insert(0, "行", df.index + 2)  # Excel上の行番号。算出根拠に残す
    return df


def metric(name: str, r: pd.Series, capacity) -> tuple[float | None, str]:
    """指標の値と、根拠として残す計算式を返す。"""
    if name == "稼働率":
        if pd.isna(capacity):
            return None, "定員がないため計算不可"
        return r.延べ利用者数 / (capacity * r.営業日数) * 100, \
            f"延べ利用者数{r.延べ利用者数:,.0f} ÷ (定員{capacity:.0f} × 営業日数{r.営業日数:.0f}) × 100"
    if name == "延べ利用者数":
        return float(r.延べ利用者数), f"延べ利用者数{r.延べ利用者数:,.0f}"
    if name == "1人当たり単位数":
        return r.算定単位数 / r.スタッフ数, f"算定単位数{r.算定単位数:,.0f} ÷ スタッフ数{r.スタッフ数:g}"
    sys.exit(f"[エラー] 評価ルール「基準値」: 未対応の指標です: {name}")


def evaluate(month: pd.Period, data_dir: Path, rules_path: Path):
    fac = read_csv(data_dir / "事業所マスタ.csv")
    res = read_csv(data_dir / "月次実績.csv")
    fac["開設年月"] = pd.PeriodIndex(fac["開設年月"], freq="M")
    res["年月"] = pd.PeriodIndex(res["年月"], freq="M")
    actual = res.set_index(["事業所ID", "年月"])

    settings = dict(zip(*read_rule_sheet(rules_path, "設定")[["項目", "値"]].T.values))
    n_months = int(settings["評価期間(月数)"])
    new_months = int(settings["新規判定(開設後の月数)"])
    cap = float(settings["達成率の上限(%)"])
    base = read_rule_sheet(rules_path, "基準値")
    ranks = read_rule_sheet(rules_path, "評価ランク").sort_values("スコア下限", ascending=False)
    exc = read_rule_sheet(rules_path, "例外条件")
    exc = exc[exc["有効"] == "○"]
    period = list(pd.period_range(end=month, periods=n_months, freq="M"))

    def rank_of(score):
        if pd.isna(score):
            return ""
        return next(r.ランク for r in ranks.itertuples() if score >= r.スコア下限)

    detail, scores, review = [], {}, []
    for f in fac.itertuples():
        for m in period:
            if m < f.開設年月 or (f.事業所ID, m) not in actual.index:
                continue
            kind = "新規" if (m - f.開設年月).n < new_months else "既存"
            rows = base[(base["事業種別"] == f.事業種別) & (base["開設区分"] == kind)]
            if rows.empty:
                review.append((f.事業所ID, f.事業所名, str(m), "ルール不足", f"基準値表に「{f.事業種別}/{kind}」がありません"))
                continue
            weighted = total = 0.0
            for b in rows.itertuples():
                value, formula = metric(b.指標, actual.loc[(f.事業所ID, m)], f.定員)
                ach = min(value / b.基準値 * 100, cap)
                weighted, total = weighted + ach * b.重み, total + b.重み
                detail.append({"事業所ID": f.事業所ID, "事業所名": f.事業所名, "年月": str(m), "開設区分": kind,
                               "指標": b.指標, "実績値": value, "計算式": formula, "基準値": b.基準値,
                               "基準値表の行": b.行, "達成率(%)": ach, "重み": b.重み})
            scores[(f.事業所ID, m)] = round(weighted / total, 1)

    avg_col = f"{n_months}ヶ月平均スコア"
    out = []
    for f in fac.itertuples():
        own = {m: scores[(f.事業所ID, m)] for m in period if (f.事業所ID, m) in scores}
        judge, applied = ["通常"], []
        missing = [m for m in period if m >= f.開設年月 and (f.事業所ID, m) not in actual.index]
        for e in exc.itertuples():
            msg = None
            if e.条件 == "開設経過月数_未満" and (month - f.開設年月).n + 1 < e.閾値:
                msg = f"開設{(month - f.開設年月).n + 1}ヶ月"
            elif e.条件 == "欠損月数_以上" and len(missing) >= e.閾値:
                msg = f"実績が未提出の月: {', '.join(map(str, missing))}"
            elif e.条件 == "稼働率_超過":
                over = [(d["年月"], d["実績値"]) for d in detail
                        if d["事業所ID"] == f.事業所ID and d["指標"] == "稼働率" and d["実績値"] > e.閾値]
                if over:
                    msg = "稼働率が" + f"{e.閾値:g}%を超える月: " + ", ".join(f"{m}({v:.0f}%)" for m, v in over)
            if msg:
                judge.append(e.処理)
                applied.append(f"{e.ルールID}:{e.説明}")
                if e.処理 == "要確認":
                    review.append((f.事業所ID, f.事業所名, "", e.ルールID, msg))
        cur = own.get(month)
        avg = round(sum(own.values()) / len(own), 1) if own else None
        out.append({"事業所ID": f.事業所ID, "事業所名": f.事業所名, "事業種別": f.事業種別,
                    "当月スコア": cur, "当月ランク": rank_of(cur), avg_col: avg, "平均ランク": rank_of(avg),
                    "評価月数": len(own), "判定": max(judge, key=PRIORITY.get), "適用ルール": " / ".join(applied)})
    result = pd.DataFrame(out)
    normal = result["判定"] == "通常"
    result["当月順位"] = result.loc[normal, "当月スコア"].rank(ascending=False, method="min").astype("Int64")
    review_df = pd.DataFrame(review, columns=["事業所ID", "事業所名", "年月", "ルールID", "内容"])
    review_df[["確認者", "確認結果", "確認日"]] = ""  # 人が記入する承認欄
    return result, pd.DataFrame(detail), review_df


def write_excel(path: Path, result, detail, review) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        result.to_excel(w, sheet_name="事業所別評価", index=False)
        detail.to_excel(w, sheet_name="算出根拠", index=False)
        review.to_excel(w, sheet_name="要確認一覧", index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        heads = [c.value for c in ws[1]]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                h = heads[c.column - 1]
                if isinstance(c.value, float):
                    c.number_format = "#,##0" if abs(c.value) >= 1000 else "0.0"
                if "ランク" in str(h) and c.value in RANK_FILL:
                    c.fill = PatternFill("solid", fgColor=RANK_FILL[c.value])
                if h == "判定" and c.value in JUDGE_FILL:
                    c.fill = PatternFill("solid", fgColor=JUDGE_FILL[c.value])
        for i, col in enumerate(ws.columns, start=1):
            w_ = max(sum(2 if ord(ch) > 255 else 1 for ch in str(c.value or "")) for c in col) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(w_, 8), 60)
        ws.freeze_panes = "C2"
    wb.save(path)


def main() -> None:
    p = argparse.ArgumentParser(description="事業所評価の簡易デモ")
    p.add_argument("--month", required=True, help="評価対象月(YYYY-MM)")
    p.add_argument("--data", default=BASE / "data" / "sample", type=Path)
    p.add_argument("--rules", default=BASE / "rules" / "評価ルール.xlsx", type=Path)
    p.add_argument("--out", default=BASE / "output", type=Path)
    a = p.parse_args()
    month = pd.Period(a.month, freq="M")
    result, detail, review = evaluate(month, a.data, a.rules)
    path = a.out / f"評価結果_{month}.xlsx"
    write_excel(path, result, detail, review)
    print(f"事業所: {len(result)}件 / 要確認: {len(review)}件")
    print(f"出力: {path}")


if __name__ == "__main__":
    main()
