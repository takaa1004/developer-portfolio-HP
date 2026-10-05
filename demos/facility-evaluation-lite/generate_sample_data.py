"""デモ用の架空データと、仮の評価ルール表を作る。

実在の事業所とは関係ありません。乱数は固定なので、何度実行しても同じデータになります。

    python generate_sample_data.py
"""

import random
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).parent
MONTHS = pd.period_range("2026-04", "2026-09", freq="M")

FACILITIES = [
    # ID, 名称, 事業種別, 定員, 開設年月
    ("F01", "デイサービスさくら", "通所介護", 18, "2018-11"),
    ("F02", "デイサービスひまわり", "通所介護", 25, "2013-07"),
    ("F03", "デイサービスみどり", "通所介護", 20, "2020-05"),  # 実績が低め
    ("F04", "デイサービスつばさ", "通所介護", 30, "2016-09"),
    ("F05", "ヘルパーステーションひかり", "訪問介護", None, "2013-04"),
    ("F06", "ヘルパーステーションこもれび", "訪問介護", None, "2026-06"),  # 開設4ヶ月
    ("F07", "ヘルパーステーションあおば", "訪問介護", None, "2012-04"),
    ("F08", "ヘルパーステーションはるかぜ", "訪問介護", None, "2016-12"),  # 7月の実績が未提出
    ("F09", "デイサービスゆうひ", "通所介護", 20, "2017-11"),  # 8月に入力ミス
    ("F10", "ヘルパーステーションそよかぜ", "訪問介護", None, "2012-09"),
]
STRENGTH = {"F03": 0.78, "F06": 0.85}


def build_results(rng: random.Random) -> pd.DataFrame:
    rows = []
    for fid, _, kind, capacity, opened in FACILITIES:
        base = STRENGTH.get(fid, rng.uniform(0.92, 1.12))
        for m in MONTHS:
            if m < pd.Period(opened, freq="M"):
                continue
            if fid == "F08" and str(m) == "2026-07":
                continue
            f = base * rng.uniform(0.95, 1.05)
            days = 26 if m.month in (5, 7, 8) else 25
            if kind == "通所介護":
                staff = round(capacity * 0.3, 1)
                visits = round(capacity * days * min(0.8 * f, 0.97))
                units = round(visits * rng.uniform(830, 870))
            else:
                staff = round(rng.uniform(9, 11), 1)
                visits = round(550 * f)
                units = round(visits * rng.uniform(410, 430))
            if fid == "F09" and str(m) == "2026-08":
                visits *= 10
            rows.append((fid, str(m), visits, units, days, staff))
    return pd.DataFrame(rows, columns=["事業所ID", "年月", "延べ利用者数", "算定単位数", "営業日数", "スタッフ数"])


def build_rules() -> dict[str, pd.DataFrame]:
    return {
        "設定": pd.DataFrame([
            ("評価期間(月数)", 6, "評価対象月を含めて遡る月数"),
            ("新規判定(開設後の月数)", 12, "開設からこの月数未満は「新規」の基準値を使う"),
            ("達成率の上限(%)", 150, "1指標の達成率がこれを超えたら上限で打ち切る"),
        ], columns=["項目", "値", "説明"]),
        "基準値": pd.DataFrame([
            ("通所介護", "既存", "稼働率", 80, 0.6), ("通所介護", "既存", "1人当たり単位数", 58000, 0.4),
            ("通所介護", "新規", "稼働率", 60, 0.6), ("通所介護", "新規", "1人当たり単位数", 45000, 0.4),
            ("訪問介護", "既存", "延べ利用者数", 550, 0.4), ("訪問介護", "既存", "1人当たり単位数", 23000, 0.6),
            ("訪問介護", "新規", "延べ利用者数", 400, 0.4), ("訪問介護", "新規", "1人当たり単位数", 17000, 0.6),
        ], columns=["事業種別", "開設区分", "指標", "基準値", "重み"]),
        "評価ランク": pd.DataFrame(
            [("S", 110), ("A", 100), ("B", 90), ("C", 80), ("D", 0)], columns=["ランク", "スコア下限"]),
        "例外条件": pd.DataFrame([
            ("EX01", "開設経過月数_未満", 6, "参考評価", "開設6ヶ月未満は参考評価", "○"),
            ("EX02", "欠損月数_以上", 1, "要確認", "評価期間内に実績が未提出の月がある", "○"),
            ("EX03", "稼働率_超過", 100, "要確認", "稼働率が100%を超える月がある(入力ミスの可能性)", "○"),
        ], columns=["ルールID", "条件", "閾値", "処理", "説明", "有効"]),
    }


def style_rules(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max(sum(2 if ord(ch) > 255 else 1 for ch in str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 60)

    def choices(ws, col: str, options: list[str]) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}100")
        ws.add_data_validation(dv)

    choices(wb["基準値"], "B", ["既存", "新規"])
    choices(wb["基準値"], "C", ["稼働率", "延べ利用者数", "1人当たり単位数"])
    choices(wb["例外条件"], "B", ["開設経過月数_未満", "欠損月数_以上", "稼働率_超過"])
    choices(wb["例外条件"], "D", ["要確認", "参考評価"])
    choices(wb["例外条件"], "F", ["○", "×"])
    wb.save(path)


def main() -> None:
    rng = random.Random(20261005)
    (BASE / "data" / "sample").mkdir(parents=True, exist_ok=True)
    (BASE / "rules").mkdir(exist_ok=True)
    facilities = pd.DataFrame(FACILITIES, columns=["事業所ID", "事業所名", "事業種別", "定員", "開設年月"])
    # 社内システムからの出力を想定して Shift-JIS で保存する
    facilities.to_csv(BASE / "data" / "sample" / "事業所マスタ.csv", index=False, encoding="cp932")
    build_results(rng).to_csv(BASE / "data" / "sample" / "月次実績.csv", index=False, encoding="cp932")
    path = BASE / "rules" / "評価ルール.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in build_rules().items():
            df.to_excel(w, sheet_name=name, index=False)
    style_rules(path)
    print("サンプルデータとルール表を作成しました")


if __name__ == "__main__":
    main()
